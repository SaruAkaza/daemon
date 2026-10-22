from __future__ import annotations

from functools import partial
import hashlib
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
import threading

from docx import Document
from PIL import Image, ImageDraw, ImageFont
from playwright.sync_api import expect
import pytest
try:
    import pymupdf
except ImportError:
    import fitz as pymupdf


@pytest.fixture(scope="module")
def library_url():
    dist = Path(__file__).resolve().parents[2] / "web/library/dist"
    if not (dist / "assets/library.js").exists():
        pytest.fail("Build the library first: cd web/library && npm ci && npm run build")
    class QuietHandler(SimpleHTTPRequestHandler):
        def log_message(self, *_args):
            pass
    server = ThreadingHTTPServer(("127.0.0.1", 0), partial(QuietHandler, directory=str(dist)))
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield f"http://127.0.0.1:{server.server_port}"
    finally:
        server.shutdown()
        thread.join(timeout=5)
        server.server_close()


@pytest.fixture
def library_files(tmp_path):
    document = Document()
    document.add_heading("Poderes", 1)
    document.add_heading("Voz da Sentinela", 2)
    document.add_paragraph("Dano: 1d6; custo: lO; ﬁm. Texto sintético.")
    document.add_paragraph("<script>window.pwned=true</script>")
    document.add_paragraph("Equipamentos")  # Ordinary Word style; explicit category title.
    document.add_paragraph("Faca de teste. Dano: 1d4.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Tabela"
    table.cell(0, 1).text = "20%"
    document.sections[0].header.paragraphs[0].text = "Cabeçalho preservado"
    docx = tmp_path / "Livro Word.docx"
    document.save(docx)
    pdf_path = tmp_path / "Livro digital.pdf"
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        page.insert_text((50, 50), "REGRAS DA SENTINELA", fontsize=24)
        page.insert_textbox(pymupdf.Rect(50, 100, 550, 230), "Texto digital sintetico. Dano: 1d6. Custo: 2 pontos. Uma sentinela vigia os portoes da cidade.", fontsize=14)
        pdf.new_page()
        pdf.save(pdf_path)
    font = ImageFont.truetype("DejaVuSans.ttf", 48)
    images = []
    for number in (2, 10):
        image = Image.new("RGB", (1400, 600), "white")
        ImageDraw.Draw(image).multiline_text((65, 60), f"RITUAIS DA SENTINELA\nImagem {number}. Dano: 1d6.\nCusto: 2 pontos.\nUma luz protege a cidade.", font=font, fill="black", spacing=25)
        path = tmp_path / f"pagina-{number}.png"
        image.save(path)
        images.append(path)
    scanned = tmp_path / "Livro escaneado.pdf"
    with pymupdf.open() as pdf:
        page = pdf.new_page(width=700, height=300)
        page.insert_image(page.rect, filename=str(images[0]))
        pdf.save(scanned)
    return {"word": docx, "pdf": pdf_path, "images": images, "scanned": scanned}


def stored_summary(page):
    return page.evaluate("""async () => {
      const db = await new Promise((resolve,reject) => { const r=indexedDB.open('daemon-personal-library',1); r.onsuccess=()=>resolve(r.result); r.onerror=()=>reject(r.error); });
      const books=await new Promise((resolve,reject)=>{ const r=db.transaction('books').objectStore('books').getAll(); r.onsuccess=()=>resolve(r.result); r.onerror=()=>reject(r.error); });
      return books.map(b=>({id:b.id,title:b.title,pageBasis:b.pageBasis,units:b.units,sourceNames:b.originals.map(o=>o.name)}));
    }""")


def test_upload_read_search_persist_and_restore_originals(page, library_url, library_files, tmp_path):
    errors = []
    page.on("pageerror", lambda error: errors.append(str(error)))
    page.goto(library_url)
    page.get_by_label("Arquivos dos livros", exact=True).set_input_files([str(library_files["word"]), str(library_files["pdf"])])
    expect(page.locator("#book-count")).to_have_text("2", timeout=30000)
    expect(page.locator("#conversion")).to_be_hidden()
    summary = stored_summary(page)
    word = next(b for b in summary if b["pageBasis"] == "DOCX_UNVERIFIED")
    assert all(unit["page"] is None for unit in word["units"])
    assert any("Cabeçalho preservado" in unit["rawText"] for unit in word["units"])
    assert any(block["rawText"] == "Equipamentos" and block["heading"] for unit in word["units"] for block in unit["blocks"])
    pdf = next(b for b in summary if b["pageBasis"] == "PDF_PHYSICAL_INDEX")
    assert [u["page"] for u in pdf["units"]] == [1, 2]
    assert pdf["units"][1]["method"] == "blank"
    page.get_by_role("button", name="Ver página original", exact=True).click()
    expect(page.locator("#source-canvas")).to_be_visible()
    page.get_by_role("button", name="Fechar fonte original").click()
    page.get_by_label("Livro", exact=True).select_option(word["id"])
    page.get_by_label("Buscar nos livros", exact=True).fill("custo lO")
    expect(page.locator("#reader")).to_contain_text("custo: lO")
    assert page.evaluate("window.pwned") is None
    page.get_by_label("Buscar nos livros", exact=True).fill("")
    page.get_by_label("Arquivos dos livros", exact=True).set_input_files(str(library_files["word"]))
    expect(page.locator("#messages")).to_contain_text("já está na sua biblioteca")
    expect(page.locator("#book-count")).to_have_text("2")
    with page.expect_download() as downloaded:
        page.get_by_role("button", name="Exportar biblioteca", exact=True).click()
    backup = tmp_path / "library.zip"
    downloaded.value.save_as(backup)
    page.reload()
    expect(page.locator("#book-count")).to_have_text("2")
    context = page.context.browser.new_context()
    try:
        restored = context.new_page()
        restored.goto(library_url)
        expect(restored.locator("#book-count")).to_have_text("0")
        restored.get_by_label("Cópia da biblioteca", exact=True).set_input_files(backup)
        expect(restored.locator("#book-count")).to_have_text("2", timeout=15000)
        restored.get_by_label("Livro", exact=True).select_option(word["id"])
        with restored.expect_download() as original:
            restored.get_by_role("button", name="Baixar original", exact=True).click()
        output = tmp_path / "restored.docx"
        original.value.save_as(output)
        assert hashlib.sha256(output.read_bytes()).digest() == hashlib.sha256(library_files["word"].read_bytes()).digest()
    finally:
        context.close()
    assert errors == []


def test_real_ocr_scanned_pdf_and_naturally_ordered_image_book(page, library_url, library_files):
    errors, requests = [], []
    page.on("pageerror", lambda error: errors.append(str(error)))
    page.context.on("request", lambda request: requests.append((request.method, request.url)))
    page.goto(library_url)
    page.get_by_label("Arquivos dos livros", exact=True).set_input_files([str(p) for p in reversed(library_files["images"])])
    expect(page.locator("#book-count")).to_have_text("1", timeout=90000)
    image_book = stored_summary(page)[0]
    assert image_book["sourceNames"] == ["pagina-2.png", "pagina-10.png"]
    assert [u["page"] for u in image_book["units"]] == [1, 2]
    assert all(u["method"] == "ocr" and "SENTINELA" in u["rawText"] for u in image_book["units"])
    expect(page.locator("#reader .tag")).to_have_text("Magias e rituais")
    page.get_by_role("button", name="Ver página original", exact=True).click()
    expect(page.locator("#source-canvas")).to_be_visible()
    page.get_by_role("button", name="Fechar fonte original").click()
    page.get_by_label("Arquivos dos livros", exact=True).set_input_files(str(library_files["scanned"]))
    expect(page.locator("#book-count")).to_have_text("2", timeout=90000)
    scanned = next(b for b in stored_summary(page) if b["pageBasis"] == "PDF_PHYSICAL_INDEX")
    assert scanned["units"][0]["method"] == "ocr"
    assert "SENTINELA" in scanned["units"][0]["rawText"]
    assert all(method == "GET" and (url.startswith(library_url) or url.startswith(("blob:", "data:"))) for method, url in requests)
    assert errors == []


def test_cancel_and_bad_files_do_not_create_partial_books_or_block_next_file(page, library_url, library_files, tmp_path):
    page.goto(library_url)
    page.get_by_label("Arquivos dos livros", exact=True).set_input_files(str(library_files["images"][0]))
    expect(page.locator("#conversion")).to_be_visible()
    page.get_by_role("button", name="Cancelar importação", exact=True).click()
    expect(page.locator("#conversion")).to_be_hidden(timeout=15000)
    expect(page.locator("#book-count")).to_have_text("0")
    bad = tmp_path / "arquivo.pdf"
    bad.write_text("not a PDF")
    page.get_by_label("Arquivos dos livros", exact=True).set_input_files([str(bad), str(library_files["word"])])
    expect(page.locator("#messages")).to_contain_text("Formato não reconhecido")
    expect(page.locator("#book-count")).to_have_text("1", timeout=15000)
    page.set_viewport_size({"width": 390, "height": 844})
    assert page.evaluate("document.documentElement.scrollWidth <= innerWidth")
