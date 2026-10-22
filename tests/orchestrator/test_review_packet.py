from __future__ import annotations

import base64
import hashlib
import io
import json
from pathlib import Path
import re
import subprocess
import sys
import zipfile

from docx import Document
from PIL import Image
import pytest

from scripts.orchestrator.checkpoint_manager import StageContext
from scripts.orchestrator.real_adapters import build_executors
from scripts.orchestrator.review_packet import ROOT, build_review_packet, write_packet
from tests.orchestrator.test_real_adapters import make_pdf, source_manifest


def evidence_for(path: Path, manifest: dict, tmp_path: Path) -> tuple[Path, str]:
    workspace = tmp_path / "attempt"
    workspace.mkdir(exist_ok=True)
    context = StageContext(manifest["sourceId"], "extraction", manifest, path, workspace, (), 1)
    build_executors()["extraction"].run(context)
    evidence = workspace / "execution-result.json"
    return evidence, hashlib.sha256(evidence.read_bytes()).hexdigest()


def review_fixture(tmp_path: Path, *, many=False) -> tuple[Path, dict, Path, str]:
    path = tmp_path / "original.docx"
    document = Document()
    document.add_paragraph('</script><script>window.pwned=true</script><img src="https://example.invalid/x"> @@CSS@@')
    document.add_paragraph("Texto sintético; Dano 1d6; custo lO; ﬁm.")
    document.add_paragraph("")
    image = io.BytesIO()
    Image.new("RGB", (20, 12), "green").save(image, "PNG")
    document.add_picture(io.BytesIO(image.getvalue()))
    # The same image has two actual uses, but only one package member.
    document.add_picture(io.BytesIO(image.getvalue()))
    if many:
        for index in range(42):
            document.add_paragraph(f"Trecho sintético {index}")
    document.save(path)
    manifest = source_manifest(path)
    manifest["title"] = '<script>window.pwned=true</script> @@SCRIPT@@'
    evidence, sha = evidence_for(path, manifest, tmp_path)
    return path, manifest, evidence, sha


def packet_payload(data: bytes) -> dict:
    match = re.search(rb'<script id="packet-data" type="application/json">(.*?)</script>', data, re.S)
    assert match
    return json.loads(match[1])


def test_packet_preserves_text_source_hashes_and_media_usages_without_inventing_pages(tmp_path):
    path, manifest, evidence, sha = review_fixture(tmp_path)
    before = path.read_bytes(), evidence.read_bytes()
    data = build_review_packet(path, manifest, evidence, evidence_sha256=sha)
    payload = packet_payload(data)
    assert before == (path.read_bytes(), evidence.read_bytes())
    assert payload["extraction"] == json.loads(before[1])["metadata"]["extraction"]
    assert payload["extraction"]["pageCount"] is None
    assert all(b["page"] is None for b in payload["extraction"]["blocks"])
    assert payload["evidenceSha256"] == sha and payload["status"] == "DRAFT_UNVERIFIED"
    assert payload["publicationMode"] == "NOT_PUBLIC" and payload["visibility"] == "private"
    assert len(payload["media"]) == 1 and len(payload["media"][0]["anchors"]) == 2
    item = payload["media"][0]
    with zipfile.ZipFile(path) as package:
        assert item["sha256"] == hashlib.sha256(package.read(item["part"])).hexdigest()
    raster = base64.b64decode(item["preview"].split(",")[1])
    assert hashlib.sha256(raster).hexdigest() == item["previewSha256"]
    with Image.open(io.BytesIO(raster)) as image:
        assert image.size == (20, 12)
    assert b"</script><script>window.pwned" not in data
    assert "@@CSS@@" in payload["extraction"]["blocks"][0]["rawText"]


def test_packet_rejects_modified_evidence_even_if_schema_and_new_hash_are_valid(tmp_path):
    path, manifest, evidence, sha = review_fixture(tmp_path)
    result = json.loads(evidence.read_bytes())
    result["metadata"]["extraction"]["blocks"][1]["rawText"] = "Dano 999d100"
    evidence.write_text(json.dumps(result), encoding="utf-8")
    with pytest.raises(ValueError, match="evidence SHA-256"):
        build_review_packet(path, manifest, evidence, evidence_sha256=sha)
    with pytest.raises(ValueError, match="differs from source"):
        build_review_packet(path, manifest, evidence, evidence_sha256=hashlib.sha256(evidence.read_bytes()).hexdigest())
    path.write_bytes(b"another source")
    with pytest.raises(ValueError, match="Source snapshot SHA-256"):
        build_review_packet(path, manifest, evidence, evidence_sha256=hashlib.sha256(evidence.read_bytes()).hexdigest())


def test_packet_output_is_private_exclusive_and_outside_checkout_even_via_symlink(tmp_path):
    target = tmp_path / "private" / "review.html"
    write_packet(target, b"first notes")
    assert target.stat().st_mode & 0o777 == 0o600
    with pytest.raises(FileExistsError):
        write_packet(target, b"replacement")
    assert target.read_bytes() == b"first notes"
    with pytest.raises(ValueError, match="outside"):
        write_packet(ROOT / "docs" / "review.html", b"private")
    link = tmp_path / "checkout-link"
    link.symlink_to(ROOT, target_is_directory=True)
    with pytest.raises(ValueError, match="outside"):
        write_packet(link / "data" / "review.html", b"private")


def test_packet_keeps_pdf_physical_indices_including_blank_pages(tmp_path):
    path, manifest = make_pdf(tmp_path)
    evidence, sha = evidence_for(path, manifest, tmp_path)
    payload = packet_payload(build_review_packet(path, manifest, evidence, evidence_sha256=sha))
    assert payload["extraction"]["pageBasis"] == "PDF_PHYSICAL_INDEX"
    assert [b["page"] for b in payload["extraction"]["blocks"]] == [1, 2]
    assert payload["extraction"]["blocks"][1]["rawText"] == ""
    assert payload["media"] == []


def test_media_limits_fail_closed_without_writing_output(tmp_path, monkeypatch):
    from scripts.orchestrator import review_packet
    path, manifest, evidence, sha = review_fixture(tmp_path)
    monkeypatch.setattr(review_packet, "MAX_IMAGE_PIXELS", 10)
    with pytest.raises(ValueError, match="pixel resource limit"):
        build_review_packet(path, manifest, evidence, evidence_sha256=sha)
    monkeypatch.setattr(review_packet, "MAX_MEDIA_BYTES", 1)
    with pytest.raises(ValueError, match="media resource limit"):
        build_review_packet(path, manifest, evidence, evidence_sha256=sha)


def test_external_image_references_are_reported_without_fetching_and_unsupported_media_is_retained(tmp_path):
    path, manifest, _, _ = review_fixture(tmp_path)
    with zipfile.ZipFile(path) as package:
        files = {name: package.read(name) for name in package.namelist()}
    rel = "word/_rels/document.xml.rels"
    import xml.etree.ElementTree as ET
    root = ET.fromstring(files[rel])
    for element in root:
        if element.get("Type", "").endswith("/image"):
            element.set("TargetMode", "External")
            element.set("Target", "https://example.invalid/private.png")
    files[rel] = ET.tostring(root)
    files["word/media/unsupported.svg"] = b'<svg xmlns="http://www.w3.org/2000/svg"><script>alert(1)</script></svg>'
    with zipfile.ZipFile(path, "w") as package:
        for name, raw in files.items():
            package.writestr(name, raw)
    manifest = source_manifest(path)
    evidence, sha = evidence_for(path, manifest, tmp_path)
    payload = packet_payload(build_review_packet(path, manifest, evidence, evidence_sha256=sha))
    assert len(payload["unresolvedMediaReferences"]) == 2
    assert all(not item["anchors"] for item in payload["media"])
    unsupported = next(m for m in payload["media"] if m["part"].endswith(".svg"))
    assert unsupported["preview"] is None
    assert b"<svg" not in build_review_packet(path, manifest, evidence, evidence_sha256=sha)


def test_cli_generates_json_summary_and_refuses_to_overwrite(tmp_path):
    path, manifest, evidence, sha = review_fixture(tmp_path)
    manifest_path = tmp_path / "manifest.json"
    manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
    output = tmp_path / "review.html"
    args = [sys.executable, "-m", "scripts.orchestrator.review_packet", "--source-root", str(tmp_path),
            "--manifest", str(manifest_path), "--evidence", str(evidence), "--evidence-sha256", sha,
            "--output", str(output)]
    run = subprocess.run(args, cwd=ROOT, capture_output=True, text=True)
    assert run.returncode == 0, run.stderr
    summary = json.loads(run.stdout)
    assert summary["sha256"] == hashlib.sha256(output.read_bytes()).hexdigest()
    again = subprocess.run(args, cwd=ROOT, capture_output=True, text=True)
    assert again.returncode != 0
