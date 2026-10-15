from __future__ import annotations

from dataclasses import replace
import hashlib
import io
import json
from pathlib import Path
import subprocess
import sys
import urllib.parse
import urllib.request
import urllib.error
import zipfile

from docx import Document
from docx.oxml import OxmlElement
import fitz
import pytest

from scripts.agents.canonical_json import canonical_json_bytes
from scripts.agents.frontend_projector import FrontendProjector
from scripts.agents.preview_server import PreviewServer
from scripts.orchestrator.batch_orchestrator import BatchOrchestrator
from scripts.orchestrator.checkpoint_manager import StageContext, digest
from scripts.orchestrator.real_adapters import build_executors
from scripts.orchestrator.source_extraction import extract_source, normalize_text
from tests.orchestrator.helpers import command, pump


def source_manifest(path: Path) -> dict:
    return {"schemaVersion": "1.0", "sourceId": "physical-fixture", "title": "Autoria sintética",
            "path": path.name, "extension": path.suffix, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "rightsStatus": "UNKNOWN", "visibility": "private", "publicationMode": "NOT_PUBLIC"}


def make_docx(tmp_path: Path) -> tuple[Path, dict]:
    path = tmp_path / "source.docx"
    document = Document()
    paragraph = document.add_paragraph("ﬁm; ﬅ; Dano 1d6; custo lO; ex-\nmarido")
    paragraph.add_run()._r.append(OxmlElement('w:lastRenderedPageBreak'))
    document.add_page_break()
    document.add_paragraph("")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Tabela"
    table.cell(0, 1).text = "10%"
    table.cell(1, 0).text = ""
    table.cell(1, 1).text = "1d8"
    document.sections[0].header.paragraphs[0].text = "Cabeçalho"
    document.sections[0].footer.paragraphs[0].text = "Rodapé"
    document.save(path)
    return path, source_manifest(path)


def make_pdf(tmp_path: Path, *, visual=False) -> tuple[Path, dict]:
    path = tmp_path / "source.pdf"
    with fitz.open() as document:
        page = document.new_page()
        page.insert_text((50, 50), "Texto privado sintetico; Dano 1d6; custo 2.")
        if visual:
            page.draw_rect(fitz.Rect(30, 70, 250, 200))
        document.new_page()  # Blank pages are retained, never dropped.
        document.save(path)
    return path, source_manifest(path)


def build_engine(tmp_path: Path, adapters=None):
    return BatchOrchestrator(tmp_path / "runtime", source_root=tmp_path,
                             executors=adapters or build_executors(), concurrency=1)


def test_docx_keeps_all_story_locations_blank_paragraphs_and_tables_without_inventing_pages(tmp_path):
    path, manifest = make_docx(tmp_path)
    before = path.read_bytes()
    audit = extract_source(path, manifest)
    blocks = audit["blocks"]
    assert path.read_bytes() == before
    assert audit["pageCount"] is None and audit["pageBasis"] == "UNVERIFIED"
    assert all(block["page"] is None for block in blocks)
    assert audit["pageHints"] == {"explicitBreakCount": 1, "renderedBreakCount": 1}
    assert audit["coverage"]["wordTextNodeCount"] == audit["coverage"]["extractedWordTextNodeCount"]
    assert audit["tableCount"] == 1
    assert len({block["anchor"] for block in blocks}) == len(blocks)
    assert any(block["rawText"] == "" for block in blocks)
    assert [b["rawText"] for b in blocks if b["kind"] == "table_paragraph"] == ["Tabela", "10%", "", "1d8"]
    assert {"Cabeçalho", "Rodapé"}.issubset({b["rawText"] for b in blocks})
    first = blocks[0]
    assert first["rawText"] == "ﬁm; ﬅ; Dano 1d6; custo lO; ex-\nmarido"
    assert first["cleanText"] == "fim; st; Dano 1d6; custo lO; ex-\nmarido"
    assert audit["issueCodes"] == ["SOURCE_PAGE_MAPPING_REQUIRED"]


@pytest.mark.parametrize("raw", ["Dano: 1d10O", "l PV e lO pontos", "ex-\nmarido", "T e n d ê n c i a", "a\u00a0b"])
def test_normalization_does_not_guess_rules_ocr_or_hyphenation(raw):
    assert normalize_text(raw) == (raw, {})


def test_real_docx_gate_creates_hash_bound_review_evidence_without_checkpoint(tmp_path):
    path, manifest = make_docx(tmp_path)
    with build_engine(tmp_path) as engine:
        engine.register("physical-fixture", manifest)
        book = engine.state["books"]["physical-fixture"]
        pump(engine, lambda: book["status"] == "NEEDS_HUMAN_REVIEW")
        review = engine.reviews.get_review(book["reviewId"])
        failed = review["evidence"]["failedAttempt"]
        assert failed["stage"] == "extraction" and failed["bookId"] == "physical-fixture"
        assert "SOURCE_PAGE_MAPPING_REQUIRED" in review["problem"]
        assert engine.checkpoints.load("physical-fixture", "extraction") is None
        assert engine.checkpoints.load("physical-fixture", "editorial") is None
        relative = next(iter(failed["artifactRefs"]))
        raw = engine.review_artifact(review["reviewId"], relative)
        assert hashlib.sha256(raw).hexdigest() == failed["artifactRefs"][relative]["sha256"]
        result = json.loads(raw)
        assert result["evidence"] == []  # There are no certified page citations to fabricate.
        assert result["status"] == "VALIDATION_FAILED"
        assert all(b["page"] is None for b in result["metadata"]["extraction"]["blocks"])
        assert str(engine.root / relative) in engine.jobs.load(book["jobId"])["artifacts"]
        old_id = book["reviewId"]
    with build_engine(tmp_path) as recovered:
        book = recovered.state["books"]["physical-fixture"]
        assert book["status"] == "NEEDS_HUMAN_REVIEW" and book["reviewId"] == old_id
        recovered.reviews.record_decision({"reviewId": old_id, "decidedBy": "human-fixture",
                                           "decision": "APPROVE", "snapshotSha256": review["snapshotSha256"]})
        recovered.tick()
        assert book["status"] == "NEEDS_HUMAN_REVIEW"  # An acknowledgement cannot supply missing pages.
        assert recovered.checkpoints.load("physical-fixture", "extraction") is None


def test_plain_pdf_has_real_physical_pages_then_stops_for_editorial_classification(tmp_path):
    path, manifest = make_pdf(tmp_path)
    with build_engine(tmp_path) as engine:
        engine.register("physical-fixture", manifest)
        book = engine.state["books"]["physical-fixture"]
        pump(engine, lambda: book["status"] == "NEEDS_HUMAN_REVIEW")
        assert book["currentStage"] == "editorial"
        cp = engine.checkpoints.load("physical-fixture", "extraction")
        result = json.loads((engine.root / next(iter(cp["artifactRefs"]))).read_bytes())
        audit = result["metadata"]["extraction"]
        assert audit["pageBasis"] == "PDF_PHYSICAL_INDEX" and audit["pageCount"] == 2
        assert audit["coverage"]["physicalPages"] == [1, 2]
        assert audit["blocks"][1]["rawText"] == ""
        review = engine.reviews.get_review(book["reviewId"])
        failed = review["evidence"]["failedAttempt"]
        editorial = json.loads(engine.review_artifact(review["reviewId"], next(iter(failed["artifactRefs"]))))
        assert [b["destination"] for b in editorial["metadata"]["editorial"]["blocks"]] == ["PENDING_CLASSIFICATION", "BLANK_PAGE"]
        assert engine.checkpoints.load("physical-fixture", "editorial") is None
        assert engine.checkpoints.load("physical-fixture", "entities") is None


def test_visual_pdf_does_not_silently_certify_only_the_text_layer(tmp_path):
    path, manifest = make_pdf(tmp_path, visual=True)
    audit = extract_source(path, manifest)
    assert audit["issueCodes"] == ["VISUAL_CONTENT_REQUIRES_REVIEW"]
    assert audit["visualPages"] == [1]
    with build_engine(tmp_path) as engine:
        engine.register("physical-fixture", manifest)
        pump(engine, lambda: engine.state["books"]["physical-fixture"]["status"] == "NEEDS_HUMAN_REVIEW")
        assert engine.checkpoints.load("physical-fixture", "extraction") is None


def test_gate_recomputes_against_source_not_only_artifact_schema(tmp_path):
    path, manifest = make_pdf(tmp_path)
    adapters = build_executors()
    original = adapters["extraction"]

    def altered(context):
        outputs = original.run(context)
        artifact = context.workspace / "execution-result.json"
        result = json.loads(artifact.read_bytes())
        result["metadata"]["extraction"]["blocks"][0]["cleanText"] = "Invented Dano 1d20"
        artifact.write_bytes(canonical_json_bytes(result))
        return outputs

    adapters["extraction"] = replace(original, run=altered)
    with build_engine(tmp_path, adapters) as engine:
        engine.register("physical-fixture", manifest)
        book = engine.state["books"]["physical-fixture"]
        pump(engine, lambda: book["status"] == "BLOCKED")
        assert "differs" in book["failure"]
        assert book["reviewId"] is None
        assert engine.checkpoints.load("physical-fixture", "extraction") is None


@pytest.mark.parametrize("change", ["bytes", "contract"])
def test_pending_review_rejects_modified_artifacts_or_contract_on_recovery(tmp_path, change):
    path, manifest = make_docx(tmp_path)
    with build_engine(tmp_path) as engine:
        engine.register("physical-fixture", manifest)
        book = engine.state["books"]["physical-fixture"]
        pump(engine, lambda: book["status"] == "NEEDS_HUMAN_REVIEW")
        old_id = book["reviewId"]
        failed = book["failedAttempt"]
        if change == "bytes":
            (engine.root / next(iter(failed["artifactRefs"]))).write_text("{}")
            with pytest.raises(ValueError):
                engine.validate_review(old_id)
    adapters = build_executors()
    if change == "contract":
        adapters["extraction"] = replace(adapters["extraction"], version="new-contract")
    with build_engine(tmp_path, adapters) as recovered:
        book = recovered.state["books"]["physical-fixture"]
        assert book["status"] == "QUEUED" and book["epochs"]["extraction"] == 1
        assert recovered.reviews.get_review(old_id)["status"] == "SUPERSEDED"
        pump(recovered, lambda: book["status"] == "NEEDS_HUMAN_REVIEW")
        assert book["reviewId"] != old_id


def test_unsafe_or_invalid_source_packages_fail_before_extraction(tmp_path, monkeypatch):
    path, manifest = make_docx(tmp_path)
    bad_hash = {**manifest, "sha256": "0" * 64}
    with pytest.raises(ValueError, match="SHA-256"):
        extract_source(path, bad_hash)
    monkeypatch.setattr("scripts.orchestrator.source_extraction.MAX_XML_BYTES", 1)
    with pytest.raises(ValueError, match="resource limit"):
        extract_source(path, manifest)


def test_nested_docx_text_fields_and_deleted_text_are_preserved_as_review_evidence(tmp_path):
    path, manifest = make_docx(tmp_path)
    with zipfile.ZipFile(path) as original:
        files = {name: original.read(name) for name in original.namelist()}
    from xml.etree import ElementTree as ET
    w = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    root = ET.fromstring(files['word/document.xml'])
    paragraph = root.find('.//' + w + 'p')
    run = ET.SubElement(paragraph, w + 'r')
    box = ET.SubElement(run, w + 'txbxContent')
    nested = ET.SubElement(ET.SubElement(ET.SubElement(box, w + 'p'), w + 'r'), w + 't')
    nested.text = 'Texto interno único'
    ET.SubElement(run, w + 'instrText').text = ' PAGE '
    removed = ET.SubElement(ET.SubElement(ET.SubElement(paragraph, w + 'del'), w + 'r'), w + 'delText')
    removed.text = 'Texto excluído preservado'
    files['word/document.xml'] = ET.tostring(root)
    with zipfile.ZipFile(path, 'w') as modified:
        for name, raw in files.items():
            modified.writestr(name, raw)
    audit = extract_source(path, source_manifest(path))
    assert sum(b['rawText'].count('Texto interno único') for b in audit['blocks']) == 1
    auxiliary = [t for b in audit['blocks'] for t in b['auxiliaryTokens']]
    assert {t['text'] for t in auxiliary} == {' PAGE ', 'Texto excluído preservado'}
    assert {'TRACKED_CHANGES_REQUIRE_REVIEW', 'COMPLEX_LAYOUT_REQUIRES_REVIEW'}.issubset(audit['issueCodes'])
    assert all(b['page'] is None for b in audit['blocks'])


def test_docx_dtd_is_rejected_without_resolving_external_entities(tmp_path):
    path = tmp_path / 'source.docx'
    with zipfile.ZipFile(path, 'w') as package:
        package.writestr('word/document.xml', '<!DOCTYPE x [<!ENTITY content SYSTEM "file:///etc/hosts">]><x/>')
    with pytest.raises(ValueError, match='DTD/entity'):
        extract_source(path, source_manifest(path))


def test_cli_real_registry_until_idle_returns_review_status_without_canonical_writes(tmp_path):
    path, manifest = make_docx(tmp_path)
    registry = tmp_path / "registry.json"
    registry.write_text(json.dumps([{"bookId": "physical-fixture", "manifest": manifest}]))
    process = subprocess.run([sys.executable, "-m", "scripts.orchestrator", "--registry", str(registry),
                              "--adapters", "scripts.orchestrator.real_adapters:build_executors",
                              "--source-root", str(tmp_path), "--runtime-dir", str(tmp_path / "runtime"),
                              "--until-idle"], capture_output=True, text=True, timeout=45)
    assert process.returncode == 0, process.stderr
    snapshot = json.loads(process.stdout)
    assert snapshot["books"][0]["pipelineStatus"] == "NEEDS_HUMAN_REVIEW"
    assert snapshot["books"][0]["completedStages"] == 1
    assert path.read_bytes() and not (tmp_path / "data").exists()
    manifest["title"] = "Changed revision"
    registry.write_text(json.dumps([{"bookId": "physical-fixture", "manifest": manifest}]))
    changed = subprocess.run(process.args, capture_output=True, text=True, timeout=45)
    assert changed.returncode != 0 and "Registry manifest changed" in changed.stderr


def test_private_download_is_authenticated_bound_to_current_review_and_unavailable_in_public(tmp_path):
    path, manifest = make_docx(tmp_path)
    with build_engine(tmp_path) as engine:
        engine.register("physical-fixture", manifest)
        book = engine.state["books"]["physical-fixture"]
        pump(engine, lambda: book["status"] == "NEEDS_HUMAN_REVIEW")
        review_id = book["reviewId"]
        relative = next(iter(book["failedAttempt"]["artifactRefs"]))
        server = PreviewServer(tmp_path / "runtime/preview", port=0, runtime_dir=engine.runtime_dir, orchestrator=engine)
        server.start()
        try:
            url_path = f"/api/ops/reviews/{review_id}/artifact?path=" + urllib.parse.quote(relative, safe="")
            url = f"http://127.0.0.1:{server.port}" + url_path
            with pytest.raises(urllib.error.HTTPError) as denied:
                urllib.request.urlopen(url)
            assert denied.value.code == 401
            # Tokens are kept in memory and sent as headers, never logged or added to URLs.
            auth = {"Authorization": "Bearer " + server.token_manager.get_token()}
            with urllib.request.urlopen(urllib.request.Request(url, headers=auth)) as response:
                assert response.headers["Cache-Control"] == "no-store"
                assert response.headers["Content-Disposition"].startswith("attachment")
                assert json.load(response)["metadata"]["extraction"]["pageCount"] is None
            with pytest.raises(urllib.error.HTTPError) as escaped:
                urllib.request.urlopen(urllib.request.Request(url.replace(urllib.parse.quote(relative, safe=""), "..%2Fstate.json"), headers=auth))
            assert escaped.value.code == 409
            command(engine, "physical-fixture", "REPROCESS_STAGE", {"stage": "EXTRACTION"})
            with pytest.raises(urllib.error.HTTPError) as stale:
                urllib.request.urlopen(urllib.request.Request(url, headers=auth))
            assert stale.value.code == 409
        finally:
            server.stop()
        public = PreviewServer(tmp_path / "public-preview", port=0, mode="public", runtime_dir=engine.runtime_dir)
        public.start()
        try:
            with pytest.raises(urllib.error.HTTPError) as denied:
                urllib.request.urlopen(f"http://127.0.0.1:{public.port}" + url_path)
            assert denied.value.code == 404
        finally:
            public.stop()
